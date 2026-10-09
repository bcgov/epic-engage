# Copyright © 2026 Province of British Columbia
#
# Licensed under the Apache License, Version 2.0 (the 'License');
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an 'AS IS' BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Tests for the Public/Proponent comment export.

The export carries only approved comments on free-text questions that are shown in the public
report and switched on for the export, grouped by question within their survey page.
"""
import copy
from datetime import datetime
from http import HTTPStatus
from io import BytesIO

from openpyxl import load_workbook

from met_api.constants.comment_status import Status as CommentStatus
from met_api.constants.membership_type import MembershipType
from met_api.constants.user import SYSTEM_REVIEWER
from met_api.models.report_setting import ReportSetting as ReportSettingModel
from met_api.utils.enums import ContentType
from met_api.utils.export_styles import QUESTION_BANNER_COLOUR, get_page_colours
from tests.utilities.factory_scenarios import TestJwtClaims
from tests.utilities.factory_utils import (
    factory_auth_header, factory_comment_model, factory_membership_model, factory_participant_model,
    factory_staff_user_model, factory_submission_model, factory_survey_and_eng_model,
    factory_survey_report_setting_model)


AWAITING_RESPONSE = '- awaiting response -'
NO_COMMENTS = 'No approved comments yet'

FORM_JSON = {
    'display': 'wizard',
    'components': [
        {
            'title': 'Demographics', 'key': 'page1', 'type': 'panel', 'components': [
                {
                    'key': 'age', 'type': 'simpleradios', 'label': 'What is your age?', 'input': True,
                    'values': [{'value': 'a1', 'label': '18-34'}],
                },
                {'key': 'connection', 'type': 'simpletextarea', 'label': 'Your connection?', 'input': True},
            ],
        },
        {
            'title': 'Hidden page', 'key': 'page2', 'type': 'panel', 'components': [
                {'key': 'hidden', 'type': 'simpletextfield', 'label': 'Hidden in report', 'input': True},
                {'key': 'not_exported', 'type': 'simpletextarea', 'label': 'Not exported', 'input': True},
            ],
        },
        {
            'title': 'Project design', 'key': 'page3', 'type': 'panel', 'components': [
                {'key': 'design', 'type': 'simpletextarea', 'label': 'Design comments?', 'input': True},
            ],
        },
    ],
}


def _survey_with_settings():
    survey, eng = factory_survey_and_eng_model()
    survey.form_json = FORM_JSON
    survey.save()
    settings = {
        'age': {'question_type': 'simpleradios'},
        'connection': {'description': 'Report description', 'export_description': 'Export description'},
        'hidden': {'display': False},
        'not_exported': {'export_display': False},
        'design': {'description': 'Report description only'},
    }
    for key, overrides in settings.items():
        setting = factory_survey_report_setting_model({
            'survey_id': survey.id,
            'question_id': key,
            'question_key': key,
            'question_type': overrides.get('question_type', 'simpletextarea'),
            'question': key,
            'display': overrides.get('display', True),
            'description': overrides.get('description'),
        })
        setting.export_display = overrides.get('export_display', True)
        setting.export_description = overrides.get('export_description')
        setting.save()
    return survey, eng


def _add_comment(survey, eng, component_id, text, *,  # pylint:disable=too-many-arguments
                 status=CommentStatus.Approved.value, reviewed_by='Staff'):
    participant = factory_participant_model()
    submission = factory_submission_model(survey.id, eng.id, participant.id, {
        'submission_json': {component_id: text},
        'created_date': datetime.now().strftime('%Y-%m-%d'),
        'comment_status_id': status,
        'reviewed_by': reviewed_by,
    })
    factory_comment_model(survey.id, submission.id, {
        'component_id': component_id,
        'text': text,
        'submission_date': datetime.now().strftime('%Y-%m-%d'),
    })


def _export(client, jwt, survey_id, claims=TestJwtClaims.staff_admin_role):
    headers = factory_auth_header(jwt=jwt, claims=claims)
    return client.get(f'/api/comments/survey/{survey_id}/sheet/proponent', headers=headers,
                      content_type=ContentType.JSON.value)


def _rows(rv):
    sheet = load_workbook(BytesIO(rv.data)).active
    return [(row[0], row[1]) for row in sheet.iter_rows(values_only=True)]


def test_proponent_export_groups_comments_by_page_and_question(client, jwt, session):  # pylint:disable=unused-argument
    """Assert the export lists approved comments under their question, within their page."""
    survey, eng = _survey_with_settings()
    _add_comment(survey, eng, 'connection', 'First approved')
    _add_comment(survey, eng, 'connection', 'Second approved')
    _add_comment(survey, eng, 'connection', 'Still pending', status=CommentStatus.Pending.value)
    _add_comment(survey, eng, 'connection', 'Rejected', status=CommentStatus.Rejected.value)
    _add_comment(survey, eng, 'connection', 'System approved', reviewed_by=SYSTEM_REVIEWER)
    _add_comment(survey, eng, 'hidden', 'Hidden in the public report')
    _add_comment(survey, eng, 'not_exported', 'Switched off for the export')

    rv = _export(client, jwt, survey.id)

    assert rv.status_code == HTTPStatus.OK
    assert rv.headers['content-type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'Public Proponent Export' in rv.headers['content-disposition']
    assert _rows(rv) == [
        (f'{survey.name} - Public / Proponent Export', None),
        ('Public Comment', 'Proponent Response'),
        ('PAGE 1 - DEMOGRAPHICS', None),
        ('Your connection?', None),
        ('Export description', None),
        ('First approved', AWAITING_RESPONSE),
        ('Second approved', AWAITING_RESPONSE),
        # Page 2 holds nothing exportable, so it is left out; page 3 keeps its own number.
        ('PAGE 3 - PROJECT DESIGN', None),
        ('Design comments?', None),
        # No export description of its own, so it falls back to the report's.
        ('Report description only', None),
        (NO_COMMENTS, None),
    ]


def test_proponent_export_for_assigned_team_member(client, jwt, session):  # pylint:disable=unused-argument
    """Assert a team member assigned to the engagement can export its comments."""
    survey, eng = _survey_with_settings()
    user = factory_staff_user_model(TestJwtClaims.team_member_role.get('sub'))
    factory_membership_model(user_id=user.id, engagement_id=eng.id, member_type=MembershipType.TEAM_MEMBER.name)

    rv = _export(client, jwt, survey.id, claims=TestJwtClaims.team_member_role)

    assert rv.status_code == HTTPStatus.OK


def test_proponent_export_forbidden_for_unassigned_team_member(client, jwt, session):  # pylint:disable=unused-argument
    """Assert a team member not assigned to the engagement cannot export its comments."""
    survey, _ = _survey_with_settings()
    factory_staff_user_model(TestJwtClaims.team_member_role.get('sub'))

    rv = _export(client, jwt, survey.id, claims=TestJwtClaims.team_member_role)

    assert rv.status_code == HTTPStatus.FORBIDDEN


def test_proponent_export_forbidden_for_unassigned_team_member_with_export_all_to_csv(
        client, jwt, session):  # pylint:disable=unused-argument
    """Assert the Team Member group's export_all_to_csv role does not open other engagements' exports."""
    survey, _ = _survey_with_settings()
    factory_staff_user_model(TestJwtClaims.team_member_role.get('sub'))
    claims = copy.deepcopy(TestJwtClaims.team_member_role.value)
    claims['realm_access']['roles'].append('export_all_to_csv')

    rv = _export(client, jwt, survey.id, claims=claims)

    assert rv.status_code == HTTPStatus.FORBIDDEN


def test_proponent_export_unauthorized_for_reviewer(client, jwt, session):  # pylint:disable=unused-argument
    """Assert a reviewer, who lacks the export role, cannot export comments."""
    survey, _ = _survey_with_settings()

    rv = _export(client, jwt, survey.id, claims=TestJwtClaims.reviewer_role)

    assert rv.status_code == HTTPStatus.UNAUTHORIZED


def test_proponent_export_survey_not_found(client, jwt, session):  # pylint:disable=unused-argument
    """Assert exporting a nonexistent survey returns 404."""
    rv = _export(client, jwt, 999999999)

    assert rv.status_code == HTTPStatus.NOT_FOUND


def _sheet(client, jwt, survey_id):
    return load_workbook(BytesIO(_export(client, jwt, survey_id).data)).active


def _row_of(sheet, value):
    return next(row[0].row for row in sheet.iter_rows(max_col=1) if row[0].value == value)


def _fills(sheet, row):
    return tuple(sheet.cell(row=row, column=column).fill.fgColor.rgb[-6:] for column in (1, 2))


def test_proponent_export_fits_both_columns_on_screen(client, jwt, session):  # pylint:disable=unused-argument
    """Assert the two columns are narrow enough to sit side by side on a laptop screen."""
    survey, _ = _survey_with_settings()

    sheet = _sheet(client, jwt, survey.id)

    assert sheet.column_dimensions['A'].width == 80
    assert sheet.column_dimensions['B'].width == 50


def test_proponent_export_uses_bc_sans_10(client, jwt, session):  # pylint:disable=unused-argument
    """Assert every cell is BC Sans 10, bold only for the title, headings and banners."""
    survey, eng = _survey_with_settings()
    _add_comment(survey, eng, 'connection', 'First approved')

    sheet = _sheet(client, jwt, survey.id)
    bold = {f'{survey.name} - Public / Proponent Export', 'Public Comment', 'Proponent Response',
            'PAGE 1 - DEMOGRAPHICS', 'PAGE 3 - PROJECT DESIGN', 'Your connection?', 'Design comments?'}

    for row in sheet.iter_rows(max_col=2):
        for cell in row:
            if cell.value is None:
                continue
            assert (cell.font.name, cell.font.sz) == ('BC Sans', 10), cell.value
            assert bool(cell.font.b) == (cell.value in bold), cell.value


def test_proponent_export_colours_each_page_banner(client, jwt, session):  # pylint:disable=unused-argument
    """Assert each page separator takes its own page colour, as in the aggregated dashboard export."""
    survey, _ = _survey_with_settings()

    sheet = _sheet(client, jwt, survey.id)

    # A banner is merged across both columns, and Excel draws it in its first cell's style.
    assert _fills(sheet, _row_of(sheet, 'PAGE 1 - DEMOGRAPHICS'))[0] == get_page_colours(0).banner
    assert _fills(sheet, _row_of(sheet, 'PAGE 3 - PROJECT DESIGN'))[0] == get_page_colours(2).banner
    assert _fills(sheet, _row_of(sheet, 'Your connection?'))[0] == QUESTION_BANNER_COLOUR


def test_proponent_export_bands_comments_in_page_colours(client, jwt, session):  # pylint:disable=unused-argument
    """Assert comment rows alternate their page's two shades, restarting at each question."""
    survey, eng = _survey_with_settings()
    _add_comment(survey, eng, 'connection', 'First')
    _add_comment(survey, eng, 'connection', 'Second')
    _add_comment(survey, eng, 'connection', 'Third')
    _add_comment(survey, eng, 'design', 'On page three')

    sheet = _sheet(client, jwt, survey.id)
    page_one, page_three = get_page_colours(0), get_page_colours(2)

    assert _fills(sheet, _row_of(sheet, 'First')) == (page_one.band_light,) * 2
    assert _fills(sheet, _row_of(sheet, 'Second')) == (page_one.band_dark,) * 2
    assert _fills(sheet, _row_of(sheet, 'Third')) == (page_one.band_light,) * 2
    assert _fills(sheet, _row_of(sheet, 'On page three')) == (page_three.band_light,) * 2


def test_proponent_export_wraps_long_comments_in_own_column(client, jwt, session):  # pylint:disable=unused-argument
    """Assert a long comment wraps inside Public Comment, in a row tall enough to show all of it."""
    survey, eng = _survey_with_settings()
    long_comment = 'A long comment about the project and its effects on the watershed. ' * 12
    _add_comment(survey, eng, 'connection', 'Short')
    _add_comment(survey, eng, 'connection', long_comment)

    sheet = _sheet(client, jwt, survey.id)
    short_row, long_row = _row_of(sheet, 'Short'), _row_of(sheet, long_comment)

    assert sheet.cell(row=long_row, column=1).alignment.wrap_text
    merged_rows = {row for merged in sheet.merged_cells.ranges for row in range(merged.min_row, merged.max_row + 1)}
    assert long_row not in merged_rows
    assert sheet.cell(row=long_row, column=2).value == AWAITING_RESPONSE
    assert sheet.row_dimensions[long_row].height > (sheet.row_dimensions[short_row].height or 15) * 4


def test_proponent_export_cleared_description_stays_cleared(client, jwt, session):  # pylint:disable=unused-argument
    """Assert an export description cleared on purpose does not fall back to the report's."""
    survey, _ = _survey_with_settings()
    setting = ReportSettingModel.find_by_question_key(survey.id, 'design')
    setting.export_description = ''
    setting.save()

    rows = [value for value, _ in _rows(_export(client, jwt, survey.id))]

    assert 'Report description only' not in rows


def test_proponent_export_includes_unreviewed_approvals(client, jwt, session):  # pylint:disable=unused-argument
    """Assert an approved comment with no reviewer recorded is still exported."""
    survey, eng = _survey_with_settings()
    _add_comment(survey, eng, 'connection', 'No reviewer recorded', reviewed_by=None)

    assert ('No reviewer recorded', AWAITING_RESPONSE) in _rows(_export(client, jwt, survey.id))


def test_proponent_export_writes_comments_as_text_not_formulas(client, jwt, session):  # pylint:disable=unused-argument
    """Assert a comment starting with '=' is kept as its text rather than becoming a live formula."""
    survey, eng = _survey_with_settings()
    _add_comment(survey, eng, 'connection', '= I disagree with this project')

    sheet = load_workbook(BytesIO(_export(client, jwt, survey.id).data)).active
    cell = next(row[0] for row in sheet.iter_rows(max_col=1) if row[0].value == '= I disagree with this project')

    assert cell.data_type == 's'


def test_proponent_export_strips_characters_excel_cannot_store(client, jwt, session):  # pylint:disable=unused-argument
    """Assert a comment holding a control character is exported without it, not failing the export."""
    survey, eng = _survey_with_settings()
    _add_comment(survey, eng, 'connection', 'line one\x0bline two')

    rv = _export(client, jwt, survey.id)

    assert rv.status_code == HTTPStatus.OK
    assert ('line oneline two', AWAITING_RESPONSE) in _rows(rv)


def test_proponent_export_header_survives_non_latin1_names(client, jwt, session):  # pylint:disable=unused-argument
    """Assert an engagement name outside Latin-1 still gives a header the server can send."""
    survey, eng = _survey_with_settings()
    eng.name = 'Tŝilhqot’in – Phase 1'
    eng.save()

    rv = _export(client, jwt, survey.id)

    assert rv.status_code == HTTPStatus.OK
    disposition = rv.headers['content-disposition']
    disposition.encode('latin-1')
    assert "filename*=UTF-8''" in disposition
