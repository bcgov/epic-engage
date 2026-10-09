# Copyright © 2026 Province of British Columbia
#
# Licensed under the Apache License, Version 2.0 (the 'License');
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an 'AS IS' BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Service for the Public/Proponent comment export.

A spreadsheet of approved comments that is safe to share with the public and the project's
proponents, who answer each comment in the column beside it. Comments are grouped by question,
and questions by their survey page, in form order.

Only free-text questions shown in the public report and switched on for the export are listed.
Every such question is listed even before it has an approved comment, so the sheet can be
shared as comments come in.
"""
from __future__ import annotations

import math
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from met_api.constants.membership_type import MembershipType
from met_api.models.comment import Comment as CommentModel
from met_api.models.report_setting import ReportSetting as ReportSettingModel
from met_api.models.survey import Survey as SurveyModel
from met_api.services import authorization
from met_api.utils.datetime import utc_datetime
from met_api.utils.export_styles import (
    BODY_FONT_COLOUR, CELL_BORDER_COLOUR, MUTED_FONT_COLOUR, QUESTION_BANNER_COLOUR, get_page_colours,
    get_zebra_colour)
from met_api.utils.roles import Role
from met_api.utils.survey_export_columns import FREE_TEXT_TYPES, build_export_columns


AWAITING_RESPONSE = '- awaiting response -'
NO_COMMENTS = 'No approved comments yet'
COLUMN_HEADINGS = ('Public Comment', 'Proponent Response')

COMMENT_HEADER_COLOUR = '5A6473'
RESPONSE_HEADER_COLOUR = '006064'
HEADER_FONT_COLOUR = 'FFFFFF'

FONT_NAME = 'BC Sans'
FONT_SIZE = 10

COMMENT_COLUMN_WIDTH = 80
RESPONSE_COLUMN_WIDTH = 50
COLUMN_COUNT = 2

CHARACTERS_PER_WIDTH = 1.1
LINE_HEIGHT = 13.5
ROW_PADDING = 3

_WRAP = Alignment(horizontal='left', vertical='top', wrap_text=True)
_CELL_EDGE = Side(style='thin', color=CELL_BORDER_COLOUR)
_THIN_BORDER = Border(left=_CELL_EDGE, right=_CELL_EDGE, top=_CELL_EDGE, bottom=_CELL_EDGE)


def _font(**kwargs) -> Font:
    return Font(name=FONT_NAME, size=FONT_SIZE, **kwargs)


def _row_height(text: str, column_width: float) -> float:
    """Estimate the height a row needs to show all of its wrapped text."""
    per_line = column_width * CHARACTERS_PER_WIDTH
    lines = sum(max(1, math.ceil(len(paragraph) / per_line)) for paragraph in text.split('\n'))
    return lines * LINE_HEIGHT + ROW_PADDING


def _set_text(cell, value: str):
    """Write text as plain text, whatever it holds.

    Comments come from the public: one starting with '=' would otherwise become a live formula,
    and a control character Excel cannot store would fail the whole export.
    """
    cell.value = ILLEGAL_CHARACTERS_RE.sub('', value)
    cell.data_type = 's'


def _fill(colour: str) -> PatternFill:
    return PatternFill(start_color=colour, end_color=colour, fill_type='solid')


class ProponentExportService:  # pylint: disable=too-few-public-methods
    """Public/Proponent comment export service."""

    @classmethod
    def export_comments_to_spread_sheet(cls, survey_id) -> tuple:
        """Build the Public/Proponent comment export for a survey.

        Returns the workbook as an in-memory byte stream along with a suggested filename.
        """
        survey = SurveyModel.find_by_id(survey_id)
        if not survey:
            raise KeyError(f'Survey with id {survey_id} not found')
        one_of_roles = (
            MembershipType.TEAM_MEMBER.name,
            Role.VIEW_PRIVATE_ENGAGEMENTS.value
        )
        authorization.check_auth(one_of_roles=one_of_roles, engagement_id=survey.engagement_id)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Public Proponent Export'
        cls._write_sheet(sheet, survey)

        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        return stream, cls._build_file_name(survey)

    @classmethod
    def _write_sheet(cls, sheet, survey: SurveyModel):
        settings = {
            setting.question_key: setting
            for setting in ReportSettingModel.find_by_survey_id(survey.id)
            if setting.display and setting.export_display
        }
        comments_by_question = {}
        for comment in CommentModel.get_proponent_export_comments_by_survey_id(survey.id):
            comments_by_question.setdefault(comment.component_id, []).append(comment.text)

        cls._write_header(sheet, survey)
        current_page = None
        for column in build_export_columns(survey.form_json, FREE_TEXT_TYPES):
            setting = settings.get(column.question_key)
            if not setting:
                continue
            if column.page_index != current_page:
                current_page = column.page_index
                page_label = f'Page {column.page_index + 1}'
                if column.page_title:
                    page_label = f'{page_label} - {column.page_title}'
                cls._write_banner(sheet, page_label.upper(), get_page_colours(column.page_index).banner,
                                  _font(bold=True, color=HEADER_FONT_COLOUR))
            cls._write_question(sheet, column.page_index, column.question_label, setting,
                                comments_by_question.get(column.question_key, []))

    @staticmethod
    def _write_header(sheet, survey: SurveyModel):
        """Write the title and the column headings, and freeze them in place."""
        sheet.column_dimensions['A'].width = COMMENT_COLUMN_WIDTH
        sheet.column_dimensions['B'].width = RESPONSE_COLUMN_WIDTH
        title = sheet.cell(row=1, column=1)
        _set_text(title, f'{survey.name} - Public / Proponent Export')
        title.font = _font(bold=True)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=COLUMN_COUNT)
        for column, (heading, colour) in enumerate(zip(COLUMN_HEADINGS,
                                                       (COMMENT_HEADER_COLOUR, RESPONSE_HEADER_COLOUR)), start=1):
            cell = sheet.cell(row=2, column=column, value=heading)
            cell.fill = _fill(colour)
            cell.font = _font(bold=True, color=HEADER_FONT_COLOUR)
            cell.border = _THIN_BORDER
        sheet.freeze_panes = 'A3'

    @classmethod
    def _write_question(cls, sheet, page_index: int, label: str, setting: ReportSettingModel, comments: list):
        """Write a question, its description and its comments, or a notice that it has none yet."""
        cls._write_banner(sheet, label, QUESTION_BANNER_COLOUR, _font(bold=True, color=BODY_FONT_COLOUR))
        # An export description left unset follows the report's; one cleared on purpose is ''.
        description = setting.export_description
        if description is None:
            description = setting.description
        if description:
            cls._write_banner(sheet, description, QUESTION_BANNER_COLOUR, _font(italic=True, color=BODY_FONT_COLOUR))
        if not comments:
            cls._write_banner(sheet, NO_COMMENTS, get_zebra_colour(page_index, 0), _font(color=MUTED_FONT_COLOUR))
        # Banded afresh under each question, as the aggregated sheet does.
        for band_index, text in enumerate(comments):
            cls._write_comment_row(sheet, text, get_zebra_colour(page_index, band_index))

    @staticmethod
    def _write_banner(sheet, value: str, colour: str, font: Font):
        """Write a full-width row: a page, a question, its description or its empty notice."""
        row = sheet.max_row + 1
        cell = sheet.cell(row=row, column=1)
        _set_text(cell, value)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=COLUMN_COUNT)
        cell.font = font
        cell.alignment = _WRAP
        # merge_cells styles only the first cell; fill the rest so the bar is solid.
        for column in range(1, COLUMN_COUNT + 1):
            sheet.cell(row=row, column=column).fill = _fill(colour)
            sheet.cell(row=row, column=column).border = _THIN_BORDER
        sheet.row_dimensions[row].height = _row_height(cell.value, COMMENT_COLUMN_WIDTH + RESPONSE_COLUMN_WIDTH)

    @staticmethod
    def _write_comment_row(sheet, text: str, band: str):
        """Write a comment beside its awaiting response, wrapped within the comment column."""
        row = sheet.max_row + 1
        comment = sheet.cell(row=row, column=1)
        _set_text(comment, text)
        comment.font = _font(color=BODY_FONT_COLOUR)
        response = sheet.cell(row=row, column=2, value=AWAITING_RESPONSE)
        response.font = _font(color=MUTED_FONT_COLOUR)
        for cell in (comment, response):
            cell.fill = _fill(band)
            cell.border = _THIN_BORDER
            cell.alignment = _WRAP
        sheet.row_dimensions[row].height = _row_height(comment.value, COMMENT_COLUMN_WIDTH)

    @staticmethod
    def _build_file_name(survey: SurveyModel) -> str:
        """Build the download filename, dated in UTC. The web client builds the same name."""
        engagement_name = survey.engagement.name if survey.engagement else survey.name
        timestamp = utc_datetime().strftime('%Y-%m-%d')
        return f'{engagement_name} - Public Proponent Export - {timestamp}.xlsx'
